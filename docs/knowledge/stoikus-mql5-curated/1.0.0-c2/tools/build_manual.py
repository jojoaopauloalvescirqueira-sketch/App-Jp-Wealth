"""Build Markdown and DOCX derivatives. Requires python-docx and Pillow; no network."""
from pathlib import Path
import re,json,hashlib,sys,textwrap,subprocess,argparse
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image,ImageDraw,ImageFont
cli=argparse.ArgumentParser(description=__doc__)
cli.add_argument('--root',type=Path,default=Path(__file__).resolve().parents[1])
cli.add_argument('--font',type=Path,default=Path('/System/Library/Fonts/Supplemental/Arial.ttf'))
args=cli.parse_args()
ROOT=args.root.resolve(); OUT=ROOT/'publication';OUT.mkdir(exist_ok=True)
TITLE='Stoikus Engenharia em MQL5 e programação com IA'
intro=(ROOT/'editorial/abertura.md').read_text()
chapters=sorted((ROOT/'chapters').glob('*.md'))
bibliography=(ROOT/'editorial/referencias.md').read_text()
text=intro+'\n\n'+'\n\n'.join(p.read_text() for p in chapters)+bibliography
(ROOT/'MANUAL.md').write_text(text)
# Minimal native OMML parser for the actual display-math subset used in the manual.
def E(tag):return OxmlElement('m:'+tag)
def run(s,normal=False):
 r=E('r')
 if normal:
  pr=E('rPr');sty=E('sty');sty.set(qn('m:val'),'p');pr.append(sty);r.append(pr)
 t=E('t');t.text=s;r.append(t);return r
class Math:
 def __init__(self,s):self.s=s;self.i=0
 def group(self):
  while self.i<len(self.s) and self.s[self.i].isspace():self.i+=1
  if self.i<len(self.s) and self.s[self.i]=='{':self.i+=1;return self.seq('}')
  return [self.atom()]
 def seq(self,end=None):
  out=[]
  while self.i<len(self.s):
   if end and self.s[self.i]==end:self.i+=1;break
   if self.s[self.i] in '^_':
    op=self.s[self.i];self.i+=1;arg=self.group();base=out.pop() if out else run('');tag='sSup' if op=='^' else 'sSub'
    # Combine repeated sub/sup in either order.
    local=base.tag.rsplit('}',1)[-1]
    if (local=='sSub' and op=='^') or (local=='sSup' and op=='_'):
     n=E('sSubSup');a=base.find(qn('m:e'));n.append(a);sub=base.find(qn('m:sub'));sup=base.find(qn('m:sup'))
     if sub is None:sub=E('sub');sub.extend(arg)
     if sup is None:sup=E('sup');sup.extend(arg)
     n.append(sub);n.append(sup)
    else:
     n=E(tag);a=E('e');a.append(base);n.append(a);b=E('sup' if op=='^' else 'sub');b.extend(arg);n.append(b)
    out.append(n)
   else:out.append(self.atom())
  return out
 def atom(self):
  if self.i>=len(self.s):return run('')
  c=self.s[self.i];self.i+=1
  if c=='{':n=E('box');e=E('e');e.extend(self.seq('}'));n.append(e);return n
  if c!='\\':return run(c)
  m=re.match(r'[A-Za-z]+',self.s[self.i:])
  if not m:
   c=self.s[self.i] if self.i<len(self.s) else '';self.i+=1;return run('‖' if c=='|' else c)
  cmd=m.group();self.i+=len(cmd)
  if cmd=='frac':
   n=E('f');a=E('num');a.extend(self.group());b=E('den');b.extend(self.group());n.extend([a,b]);return n
  if cmd=='sqrt':
   n=E('rad');pr=E('radPr');h=E('degHide');h.set(qn('m:val'),'1');pr.append(h);n.append(pr);n.append(E('deg'));a=E('e');a.extend(self.group());n.append(a);return n
  if cmd=='hat':
   n=E('acc');pr=E('accPr');ch=E('chr');ch.set(qn('m:val'),'̂');pr.append(ch);n.append(pr);a=E('e');a.extend(self.group());n.append(a);return n
  if cmd in ('operatorname','mathrm','text'):
   nodes=self.group();s=''.join(''.join(x.itertext()) for x in nodes);return run(s,normal=True)
  syms={'leq':'≤','varepsilon':'ε','epsilon':'ϵ','theta':'θ','eta':'η','nabla':'∇','phi':'φ','sigma':'σ','gamma':'γ','beta':'β','mu':'μ','delta':'δ','sum':'∑','odot':'⊙','lfloor':'⌊','rfloor':'⌋','quad':'  ','qquad':'    ','left':'','right':'','max':'max','tanh':'tanh','log':'log','exp':'exp','cdot':'·','times':'×'}
  if cmd not in syms:raise ValueError('Unsupported LaTex '+cmd)
  return run(syms[cmd],normal=cmd in ('max','tanh','log','exp'))
def mathpara(doc,latex):
 p=doc.add_paragraph();p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(8);o=E('oMathPara');m=E('oMath');m.extend(Math(latex).seq());o.append(m);p._p.append(o)
def inline(p,s):
 for piece in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)',s):
  if not piece:continue
  if piece.startswith('**'):r=p.add_run(piece[2:-2]);r.bold=True
  elif piece.startswith('`'):
   content=piece[1:-1]
   if content in ['u_t','x_t','h_{t-1}','e_j']:
    o=E('oMath');o.extend(Math(content).seq());p._p.append(o)
   else:r=p.add_run(content);r.font.name='Courier New';r.font.size=Pt(9.5)
  else:p.add_run(re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1 (\2)',piece))
# Parse every node, edge, decision and label; render only reviewed topologies.
from render_diagrams import render_graph
def diagram(body,n):
 return render_graph(body,n,OUT,args.font)

doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=Inches(.75);sec.left_margin=sec.right_margin=Inches(.8)
for sty in ['Normal','Body Text','List Bullet','List Number']:
 st=doc.styles[sty];st.font.name='Arial';st.font.size=Pt(11);st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.space_after=Pt(7);st.paragraph_format.line_spacing=1.14
for sty,size in [('Title',23),('Subtitle',15),('Heading 1',18),('Heading 2',13),('Heading 3',11.5)]:
 st=doc.styles[sty];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.keep_with_next=True
for sec in doc.sections:
 h=sec.header.paragraphs[0];h.text='JP Wealth  |  Stoikus  |  Engenharia em MQL5';h.style='Caption';h.runs[0].font.color.rgb=RGBColor(0,0,0)
 p=sec.footer.paragraphs[0];p.alignment=2;r=p.add_run('Página ');fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
doc.add_paragraph(TITLE,'Title');doc.add_paragraph('Manual técnico avançado\nRacional de desenvolvimento e harness para programação com IA','Subtitle')
doc.add_paragraph('Elaborado para JP Wealth\nRevisão editorial candidata 1.0.0\n15 de outubro de 2026')
doc.add_paragraph('Base do estudo Stoikus 1.23.0  |  Cálculo 1.9.0\nHarness 0.1.0 CANDIDATE  |  Adoção não automática')
doc.add_page_break();doc.add_paragraph('Sumário','Heading 1')
# Two-pass static contents preserves verified pagination in viewers without field updates.
toc_path=OUT/'toc-pages.json';toc=json.loads(toc_path.read_text()) if toc_path.exists() else {}
for title in re.findall(r'^# (.+)$',text,re.M)[1:]:
 p=doc.add_paragraph();p.add_run(title)
 if title in toc:
  from docx.enum.text import WD_TAB_ALIGNMENT,WD_TAB_LEADER
  p.paragraph_format.tab_stops.add_tab_stop(Inches(6.75),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
  p.add_run('\t'+str(toc[title]))
doc.add_page_break();lines=text.splitlines();i=0;diagram_count=0;equations=0;code_lines=0
while i<len(lines):
 line=lines[i]
 if i==0:i+=1;continue
 if not line.strip():i+=1;continue
 if line.strip()=='$$':
  i+=1;block=[]
  while i<len(lines) and lines[i].strip()!='$$':block.append(lines[i]);i+=1
  mathpara(doc,' '.join(block));equations+=1;i+=1;continue
 if line.startswith('```'):
  lang=line[3:].strip();i+=1;block=[]
  while i<len(lines) and not lines[i].startswith('```'):block.append(lines[i]);i+=1
  if lang=='mermaid':
   diagram_count+=1;pic=diagram('\n'.join(block),diagram_count);p=doc.add_paragraph();p.add_run().add_picture(str(pic),width=Inches(6.8));doc.add_paragraph(f'Figura {diagram_count}  '+('Transições conceituais: origem, destino e condição' if diagram_count==2 else 'Relações conceituais do fluxo'),'Caption')
  else:
   doc.add_paragraph('Exemplo sintético de programação','Caption')
   for code_index,code in enumerate(block):
    # Keep original logical lines; Word may wrap visually without adding hard newlines.
    p=doc.add_paragraph();p.paragraph_format.keep_with_next=code_index<len(block)-1;p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0;r=p.add_run(code);r.font.name='Courier New';r.font.size=Pt(9.5);code_lines+=1
   doc.add_paragraph()
  i+=1;continue
 if line.startswith('#'):
  h=re.match(r'(#+)\s+(.+)',line);level=min(3,len(h[1]));title=h[2]
  if level==1:doc.add_page_break()
  doc.add_paragraph(title,'Heading '+str(level));i+=1;continue
 if re.match(r'^[-*]\s+',line):p=doc.add_paragraph(style='List Bullet');inline(p,re.sub(r'^[-*]\s+','',line));i+=1;continue
 if re.match(r'^\d+\.\s+',line):p=doc.add_paragraph(style='List Number');inline(p,re.sub(r'^\d+\.\s+','',line));i+=1;continue
 par=[line];i+=1
 while i<len(lines) and lines[i].strip() and not lines[i].startswith(('#','```','$$')):par.append(lines[i]);i+=1
 p=doc.add_paragraph();inline(p,' '.join(par))
# Native math searchable/editable; update fields when opened.
settings=doc.settings.element;upd=OxmlElement('w:updateFields');upd.set(qn('w:val'),'true');settings.append(upd)
doc.core_properties.title=TITLE
doc.core_properties.subject='Manual técnico e harness candidato; provas nativas e adoção separadas'
doc.core_properties.author='Elaborado para JP Wealth com assistência de IA'
p=OUT/'Stoikus_Engenharia_MQL5_Manual.docx';doc.save(p)
manifest={'manualMarkdownSHA256':hashlib.sha256(text.encode()).hexdigest(),'chapterFiles':[{'path':str(p.relative_to(ROOT)),'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in chapters],'wordCount':len(re.findall(r'\b\w+\b',text)),'equationCount':equations,'diagrams':diagram_count,'codeLines':code_lines,'docx':p.name,'docxSHA256':hashlib.sha256(p.read_bytes()).hexdigest(),'renderState':'PENDING'}
(ROOT/'validation/editorial-build.json').write_text(json.dumps(manifest,indent=2)+'\n');print(json.dumps(manifest,indent=2))
